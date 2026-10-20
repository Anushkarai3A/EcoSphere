"""Extraction tests: PDF / DOCX / PPTX / TXT -> text, and error cases."""

import io

import docx
import pptx
import pytest

from app.services.extractor import ExtractionError, extract_text


def _build_docx(text: str) -> bytes:
    document = docx.Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _build_pptx(lines: list[str]) -> bytes:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = lines[0]
    slide.placeholders[1].text = "\n".join(lines[1:])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _build_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({_escape_pdf_text(text)}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    lines: list[bytes] = [b"%PDF-1.4"]
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(b"".join(lines)))
        lines.append(f"{index} 0 obj\n".encode())
        lines.append(body)
        lines.append(b"\nendobj\n")
    xref_pos = len(b"".join(lines))
    lines.append(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        lines.append(f"{offset:010d} 00000 n \n".encode())
    lines.append(b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n")
    lines.append(f"{xref_pos}".encode() + b"\n%%EOF")
    return b"".join(lines)


def test_extract_txt() -> None:
    content = (
        "Big Data Analytics Notes\n"
        "Hadoop processes large datasets across distributed clusters.\n"
        "MapReduce splits work into map and reduce phases."
    ).encode("utf-8")
    result = extract_text("bda_notes.txt", content, "text/plain")
    assert result.kind == "TXT"
    assert "Hadoop processes large datasets" in result.text
    assert result.token_count > 0


def test_extract_docx() -> None:
    content = _build_docx("DBMS Normalization\nFunctional dependency is a constraint between two attributes.")
    result = extract_text("dbms.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.kind == "DOCX"
    assert "Functional dependency" in result.text


def test_extract_pptx() -> None:
    content = _build_pptx(["Cloud Computing", "Virtualization abstracts physical hardware.", "IaaS delivers compute over the internet."])
    result = extract_text("cloud.pptx", content, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert result.kind == "PPTX"
    assert "[Slide 1]" in result.text
    assert "Virtualization abstracts" in result.text


def test_extract_pdf() -> None:
    content = _build_pdf("Python programming notes: functions, lists and dictionaries.")
    result = extract_text("python.pdf", content, "application/pdf")
    assert result.kind == "PDF"
    assert "Python programming notes" in result.text


def test_extract_normalizes_whitespace() -> None:
    content = (
        b"First line of the notes.\n\n\n"
        b"   Second   line with    spaces and enough text to pass the minimum length check.\n"
    )
    result = extract_text("notes.txt", content, "text/plain")
    assert "First line of the notes." in result.text
    assert "Second line with spaces and enough text" in result.text


def test_extract_rejects_unsupported_kind() -> None:
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_text("diagram.png", b"\x89PNG\r\n", "image/png")


def test_extract_rejects_empty_document() -> None:
    content = _build_docx("   ")
    with pytest.raises(ExtractionError, match="No readable text"):
        extract_text("empty.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def test_extract_rejects_empty_pdf() -> None:
    content = _build_pdf("   ")
    with pytest.raises(ExtractionError, match="No readable text"):
        extract_text("empty.pdf", content, "application/pdf")


def test_extract_rejects_corrupt_pdf() -> None:
    with pytest.raises(ExtractionError, match="PDF could not be parsed"):
        extract_text("corrupt.pdf", b"%PDF-1.4\nnot a valid pdf", "application/pdf")
